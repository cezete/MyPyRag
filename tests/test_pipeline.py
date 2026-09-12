import hashlib
import json
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

import pytest

from mypyrag.cli import main
from mypyrag.config import Config
from mypyrag.converter import DoclingAdapter
from mypyrag.model import Manifest, MaxStage, State, reached
from mypyrag.pipeline import Pipeline
from mypyrag.storage import (
    StabilityTracker,
    atomic_json,
    load_manifest,
    safe_name,
    save_manifest,
    sha256,
)


class FakeConverter:
    version = "fake-1"

    def __init__(self):
        self.calls = 0

    def convert(self, source, destination):
        self.calls += 1
        if source.read_text(encoding="utf-8") == "fail":
            raise RuntimeError("Deliberate conversion failure")
        # Real DoclingDocument schema, without model loading or actual conversion.
        from docling_core.types.doc import DoclingDocument
        from docling_core.types.doc.labels import DocItemLabel

        document = DoclingDocument(name=source.stem)
        document.add_text(label=DocItemLabel.TEXT, text=source.read_text(encoding="utf-8"))
        atomic_json(destination, document.export_to_dict())
        DoclingDocument.load_from_json(destination)


@pytest.fixture
def config(tmp_path, monkeypatch):
    import os

    for key in os.environ:
        if key.startswith("MYPYRAG_"):
            monkeypatch.delenv(key)
    return replace(Config.load(tmp_path), file_stable_seconds=0)


@pytest.fixture
def pipeline(config):
    return Pipeline(config, FakeConverter())


def put(pipeline, name="sample.md", text="# Example\n\nHello."):
    source = pipeline.config.in_dir / name
    source.write_text(text, encoding="utf-8")
    return source


def test_config_env_precedence(config, tmp_path, monkeypatch):
    (tmp_path / ".env").write_text("MYPYRAG_POLL_INTERVAL_SECONDS=20\nMYPYRAG_IN_DIR=inbox\n")
    monkeypatch.setenv("MYPYRAG_POLL_INTERVAL_SECONDS", "7")
    loaded = Config.load(tmp_path)
    assert loaded.poll_interval_seconds == 7
    assert loaded.in_dir == tmp_path / "inbox"
    assert loaded.qdrant_vector_size == 768


@pytest.mark.parametrize(
    "key,value,match",
    [
        ("POLL_INTERVAL_SECONDS", "abc", "expected integer"),
        ("POLL_INTERVAL_SECONDS", "0", "integer range"),
        ("FILE_STABLE_SECONDS", "-1", "integer range"),
        ("MAX_STAGE", "BROKEN", "expected CONVERTED"),
        ("MAX_STAGE", "INDEXED", "not implemented"),
        ("IN_DIR", "", "must not be empty"),
        ("DOCLING_JSON_FILENAME", "../bad.json", "plain filename"),
        ("LOG_LEVEL", "mystery", "unknown logging"),
        ("OLLAMA_URL", "invalid", "HTTP"),
    ],
)
def test_invalid_config(config, tmp_path, monkeypatch, key, value, match):
    monkeypatch.setenv(f"MYPYRAG_{key}", value)
    with pytest.raises(ValueError, match=match):
        Config.load(tmp_path)


def test_nested_roots_rejected(config):
    with pytest.raises(ValueError, match="non-nested"):
        replace(config, error_dir=config.in_dir / "error")


def test_stage_order():
    assert reached(State.CONVERTED, MaxStage.CONVERTED)
    assert not reached(State.RECEIVED, MaxStage.CONVERTED)
    assert not reached(State.CONVERTED, MaxStage.CHUNKED)
    assert reached(State.DONE, MaxStage.INDEXED)


def test_chunked_stage_is_configurable(config):
    assert replace(config, max_stage=MaxStage.CHUNKED).max_stage == MaxStage.CHUNKED


def test_identity_and_name(tmp_path):
    source = tmp_path / "x.md"
    source.write_bytes(b"abc")
    assert sha256(source) == hashlib.sha256(b"abc").hexdigest()
    assert safe_name("../a:b? c", "a" * 64) == "a_b_c--aaaaaaaa"
    assert safe_name("...", "a" * 64) == "document--aaaaaaaa"
    assert ".pdf" not in safe_name(Path("manual.pdf").stem, "a" * 64)


def test_stability(tmp_path):
    source = tmp_path / "x.md"
    source.write_text("first")
    tracker = StabilityTracker(10)
    assert not tracker.ready(source, 0)
    assert not tracker.ready(source, 9)
    assert tracker.ready(source, 10)
    source.write_text("changed size")
    assert not tracker.ready(source, 11)
    assert tracker.ready(source, 21)
    tracker.prune(set())
    assert not tracker.observations


def test_e2e_and_restart(pipeline):
    from docling_core.types.doc import DoclingDocument

    source = put(pipeline)
    content = source.read_bytes()
    assert pipeline.cycle()
    [(directory, manifest)] = pipeline.manifests()
    assert not source.exists()
    assert (directory / manifest.source_relative_path).read_bytes() == content
    assert DoclingDocument.load_from_json(directory / "JSON/document.json").texts
    assert manifest.current_state == manifest.last_successful_state == State.CONVERTED
    assert list((directory / "CHUNKS").iterdir()) == []
    assert list(pipeline.config.done_dir.iterdir()) == []
    adapter = FakeConverter()
    assert Pipeline(pipeline.config, adapter).cycle()
    assert adapter.calls == 0


def test_manifest_and_transitions(pipeline):
    directory = pipeline._register(put(pipeline))
    manifest = load_manifest(directory)
    assert Manifest.from_dict(manifest.to_dict()) == manifest
    with pytest.raises(ValueError, match="Invalid transition"):
        manifest.transition(State.CONVERTED)
    manifest.transition(State.CONVERTING)
    manifest.transition(State.CONVERTED)
    save_manifest(directory, manifest)
    assert load_manifest(directory) == manifest
    with pytest.raises(ValueError):
        Manifest.from_dict({**manifest.to_dict(), "schema_version": 100})


def test_atomic_failure_preserves_previous_file(tmp_path):
    target = tmp_path / "manifest.json"
    atomic_json(target, {"old": 1})
    with (
        patch("mypyrag.storage.os.replace", side_effect=OSError("disk error")),
        pytest.raises(OSError),
    ):
        atomic_json(target, {"new": 2})
    assert json.loads(target.read_text()) == {"old": 1}
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize("location", ["in_dir", "done_dir", "error_dir"])
def test_duplicate_all_roots(pipeline, location, caplog):
    assert pipeline.process(put(pipeline))
    [(directory, _manifest)] = pipeline.manifests()
    root = getattr(pipeline.config, location)
    if directory.parent != root:
        directory.rename(root / directory.name)
    duplicate = put(pipeline, "copy.md")
    original = duplicate.read_bytes()
    assert pipeline.process(duplicate)
    assert duplicate.read_bytes() == original
    assert len(pipeline.manifests()) == 1
    assert "Duplicate" in caplog.text


def test_failure_does_not_block_next_document(pipeline):
    put(pipeline, "a.md", "fail")
    put(pipeline, "b.md", "good")
    assert not pipeline.cycle()
    manifests = {m.original_filename: (d, m) for d, m in pipeline.manifests()}
    directory, failed = manifests["a.md"]
    assert directory.parent == pipeline.config.error_dir
    assert failed.current_state == State.ERROR
    assert failed.last_successful_state == State.RECEIVED
    assert failed.attempt_count == 1
    assert failed.failed_stage == "CONVERTING"
    assert "Deliberate" in failed.last_error
    assert (directory / failed.source_relative_path).read_text() == "fail"
    assert manifests["b.md"][1].current_state == State.CONVERTED


def test_unsupported_file(pipeline):
    source = put(pipeline, "unsupported.xyz")
    assert not pipeline.process(source)
    [(directory, manifest)] = pipeline.manifests()
    assert directory.parent == pipeline.config.error_dir
    assert "Unsupported" in manifest.last_error


def test_interrupted_conversion_resumes(pipeline):
    directory = pipeline._register(put(pipeline))
    manifest = load_manifest(directory)
    manifest.transition(State.CONVERTING)
    save_manifest(directory, manifest)
    (directory / "JSON/document.json").write_text("partial")
    assert pipeline.cycle()
    assert load_manifest(directory).current_state == State.CONVERTED


@pytest.mark.parametrize("moved", [False, True])
def test_interrupted_receipt(pipeline, moved):
    source = put(pipeline)
    original = source.read_bytes()
    directory = pipeline._register(source)
    manifest = load_manifest(directory)
    atomic_json(
        directory / "receipt.json", {"raw_path": str(source), "manifest": manifest.to_dict()}
    )
    (directory / "manifest.json").unlink()
    if not moved:
        (directory / manifest.source_relative_path).rename(source)
    assert pipeline.cycle()
    assert load_manifest(directory).current_state == State.CONVERTED
    assert (directory / manifest.source_relative_path).read_bytes() == original
    assert not (directory / "receipt.json").exists()


def test_corrupt_manifest_preserved_and_other_work_resumes(pipeline):
    bad = pipeline._register(put(pipeline, "bad.md", "bad"))
    good = pipeline._register(put(pipeline, "good.md", "good"))
    (bad / "manifest.json").write_text("broken")
    assert not pipeline.cycle()
    assert (bad / "manifest.json").read_text() == "broken"
    assert load_manifest(good).current_state == State.CONVERTED


def test_name_collision_does_not_overwrite(pipeline):
    raw = put(pipeline)
    occupied = pipeline.config.in_dir / safe_name(raw.stem, sha256(raw))
    occupied.mkdir()
    (occupied / "keep").write_text("original")
    assert pipeline.process(raw)
    assert (occupied / "keep").read_text() == "original"
    assert pipeline.manifests()[0][0] != occupied


def test_source_tampering_is_error(pipeline):
    directory = pipeline._register(put(pipeline))
    manifest = load_manifest(directory)
    (directory / manifest.source_relative_path).write_text("tampered")
    assert not pipeline.cycle()
    assert "SHA-256 mismatch" in pipeline.manifests()[0][1].last_error


def test_single_writer_lock(pipeline):
    from filelock import Timeout

    other = Pipeline(pipeline.config, FakeConverter())
    with pipeline.lock, pytest.raises(Timeout):
        other.cycle()


def test_acquisition_failure_preserves_source_and_recovers(pipeline):
    raw = put(pipeline)
    original = raw.read_bytes()
    with patch.object(Path, "rename", side_effect=OSError("Temporary move failure")):
        assert not pipeline.cycle()
    assert raw.read_bytes() == original
    assert len(list(pipeline.config.in_dir.glob("*/receipt.json"))) == 1
    assert pipeline.cycle()
    [(directory, manifest)] = pipeline.manifests()
    assert manifest.current_state == State.CONVERTED
    assert (directory / manifest.source_relative_path).read_bytes() == original


def test_hash_failure_preserves_input(pipeline):
    raw = put(pipeline)
    with patch("mypyrag.pipeline.sha256", side_effect=OSError("Cannot read input")):
        assert not pipeline.cycle()
    assert raw.exists()
    assert pipeline.manifests() == []


def test_error_move_failure_is_recovered(pipeline):
    put(pipeline, text="fail")
    with patch.object(pipeline, "_move_error", side_effect=OSError("Cannot move yet")):
        assert not pipeline.cycle()
        assert not pipeline.cycle()
    [(directory, manifest)] = pipeline.manifests()
    assert directory.parent == pipeline.config.in_dir
    assert manifest.current_state == State.ERROR
    assert manifest.attempt_count == 1
    assert "Deliberate" in manifest.last_error
    assert not pipeline.cycle()
    [(directory, manifest)] = pipeline.manifests()
    assert directory.parent == pipeline.config.error_dir
    assert manifest.attempt_count == 1


def test_managed_source_cannot_be_reingested(pipeline):
    assert pipeline.process(put(pipeline))
    [(directory, manifest)] = pipeline.manifests()
    with pytest.raises(ValueError, match="managed document"):
        pipeline.process(directory / manifest.source_relative_path)


def test_hidden_input_is_ignored(pipeline):
    hidden = put(pipeline, ".upload.md")
    assert pipeline.cycle()
    assert hidden.exists()
    assert pipeline.manifests() == []


def test_mtime_change_resets_stability(tmp_path):
    import os

    source = tmp_path / "x.md"
    source.write_bytes(b"abc")
    tracker = StabilityTracker(10)
    assert not tracker.ready(source, 0)
    stat = source.stat()
    os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000))
    assert not tracker.ready(source, 10)
    assert tracker.ready(source, 20)


def test_adapter_rejects_partial_conversion(tmp_path):
    from types import SimpleNamespace

    from docling.datamodel.base_models import ConversionStatus

    adapter = DoclingAdapter()
    adapter._converter = SimpleNamespace(
        convert=lambda source: SimpleNamespace(
            status=ConversionStatus.PARTIAL_SUCCESS, errors=["incomplete"]
        )
    )
    with pytest.raises(ValueError, match="did not succeed"):
        adapter.convert(tmp_path / "x.md", tmp_path / "document.json")
    assert not (tmp_path / "document.json").exists()


def test_cli_status_and_error(config, tmp_path, capsys):
    assert main(["--base-dir", str(tmp_path), "status"]) == 0
    assert "Total: 0" in capsys.readouterr().out
    assert main(["--base-dir", str(tmp_path), "process", str(tmp_path / "missing.md")]) == 1


def test_cli_process_and_watch(config, tmp_path):
    source = tmp_path / "sample.md"
    source.write_text("hello")
    with patch("mypyrag.cli.DoclingAdapter", FakeConverter):
        assert main(["--base-dir", str(tmp_path), "process", str(source)]) == 0
        with patch("mypyrag.cli.time.sleep", side_effect=KeyboardInterrupt):
            assert main(["--base-dir", str(tmp_path), "watch"]) == 130


@pytest.mark.integration
@pytest.mark.parametrize("extension", [".md", ".markdown", ".htm", ".docx"])
def test_real_docling_markdown(config, extension):
    from docling_core.types.doc import DoclingDocument

    pipeline = Pipeline(config, DoclingAdapter())
    raw = put(pipeline, f"sample{extension}")
    if extension == ".htm":
        raw.write_text("<html><body><h1>Example</h1><p>Hello.</p></body></html>")
    elif extension == ".docx":
        from docx import Document

        document = Document()
        document.add_heading("Example")
        document.add_paragraph("Hello.")
        document.save(raw)
    original = raw.read_bytes()
    assert pipeline.cycle()
    [(directory, manifest)] = pipeline.manifests()
    document = DoclingDocument.load_from_json(directory / "JSON/document.json")
    assert "Hello" in document.export_to_markdown()
    assert manifest.current_state == State.CONVERTED
    assert (directory / manifest.source_relative_path).read_bytes() == original
